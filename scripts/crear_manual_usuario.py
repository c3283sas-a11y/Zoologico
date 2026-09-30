from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "manuales" / "Manual_de_usuario_EcoFauna.docx"
LOGO = ROOT / "public" / "assets" / "img" / "logoEcoFauna.png"
HABITAT = ROOT / "storage" / "uploads" / "habitats" / "bosque_tropical1_acdf20ce10.jpg"

GREEN = "173F2A"
MID_GREEN = "287A4A"
GOLD = "D4A94F"
PALE = "E9F5EC"
LIGHT = "F3F8F4"
INK = "203129"
MUTED = "66736B"
WHITE = "FFFFFF"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths)))
    tblW.set(qn('w:type'), 'dxa')
    ind = OxmlElement('w:tblInd')
    ind.set(qn('w:w'), '120')
    ind.set(qn('w:type'), 'dxa')
    tblPr.append(ind)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.first_child_found_in('w:tcW')
            tcW.set(qn('w:w'), str(width))
            tcW.set(qn('w:type'), 'dxa')
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_font(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic

def text_para(doc, text='', style=None, before=0, after=6, size=11, bold=False, color=INK, italic=False, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None: p.alignment = align
    r = p.add_run(text)
    set_font(r, size, bold, color, italic)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, {1:16,2:13,3:12}[level], True, GREEN if level < 3 else MID_GREEN)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text), 11)
    return p

def numbered(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text), 11)
    return p

def callout(doc, title, body):
    t = doc.add_table(rows=1, cols=1)
    set_table_geometry(t, [9360])
    set_repeat_table_header(t.rows[0])
    c = t.cell(0,0)
    set_cell_shading(c, PALE)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(title + ' '), 10.5, True, GREEN)
    set_font(p.add_run(body), 10.5, False, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def data_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    set_table_geometry(t, widths)
    hrow = t.rows[0]
    set_repeat_table_header(hrow)
    for cell, val in zip(hrow.cells, headers):
        set_cell_shading(cell, GREEN)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(val), 9.5, True, WHITE)
    for row in rows:
        cells = t.add_row().cells
        for i, (cell, val) in enumerate(zip(cells, row)):
            if len(t.rows) % 2 == 0: set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(val), 9.5, i == 0, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.85); section.bottom_margin = Inches(0.8)
section.left_margin = Inches(1); section.right_margin = Inches(1)
section.header_distance = Inches(0.49); section.footer_distance = Inches(0.49)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Calibri'; normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in [('Heading 1',16,GREEN,18,10),('Heading 2',13,GREEN,14,7),('Heading 3',12,MID_GREEN,10,5)]:
    st = styles[name]; st.font.name='Calibri'; st._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=True; st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)

# Header/footer
header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run('EcoFauna | Manual de usuario'), 9, True, MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run('EcoFauna - Gestión administrativa'), 9, False, MUTED)

# Cover
if LOGO.exists():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(22); p.paragraph_format.space_after=Pt(12)
    p.add_run().add_picture(str(LOGO), width=Inches(1.45))
text_para(doc, 'MANUAL DE USUARIO', before=10, after=4, size=13, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER)
text_para(doc, 'Administración de cuidadores, hábitats, usuarios y permisos', before=0, after=12, size=25, bold=True, color=GREEN, align=WD_ALIGN_PARAGRAPH.CENTER)
text_para(doc, 'EcoFauna', before=0, after=20, size=14, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
if HABITAT.exists():
    p = doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(12); p.add_run().add_picture(str(HABITAT), width=Inches(5.65))
text_para(doc, 'Guía práctica para personal administrador', before=4, after=3, size=11, bold=True, color=GREEN, align=WD_ALIGN_PARAGRAPH.CENTER)
text_para(doc, 'Basado en los formularios, listados, estados y acciones disponibles en la aplicación EcoFauna.', before=0, after=0, size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

heading(doc, '1. Propósito y acceso', 1)
text_para(doc, 'Este manual explica cómo operar los apartados administrativos revisados en EcoFauna: Cuidadores, Hábitats, Usuarios y Roles y permisos. Las opciones descritas están destinadas a cuentas con rol Administrador.')
callout(doc, 'Antes de iniciar.', 'Acceda con una cuenta administradora. Las pantallas muestran botones Regresar o Administración para volver al panel central, mensajes de confirmación tras guardar y avisos de error si algún dato no cumple las condiciones.')
heading(doc, 'Navegación general', 2)
data_table(doc, ['Apartado', 'Qué permite hacer'], [
    ['Cuidadores', 'Registrar y consultar el personal cuidador; gestionar su ficha, animales, horarios y tareas.'],
    ['Hábitats', 'Registrar espacios, asignar animales, editar sus datos, filtrar el catálogo y activar o inactivar registros.'],
    ['Usuarios', 'Crear y editar cuentas, cambiar roles, actualizar datos y rehabilitar cuentas bloqueadas.'],
    ['Roles y permisos', 'Definir qué módulos puede utilizar cada rol no administrativo.'],
], [2200,7160])
heading(doc, 'Convenciones de pantalla', 2)
bullet(doc, 'Los campos marcados con asterisco (*) son obligatorios.')
bullet(doc, 'Los botones de guardar registran la información; Cancelar edición descarta la edición en curso y vuelve al listado.')
bullet(doc, 'Los distintivos Activo/Inactivo o Activa/Bloqueada reflejan el estado actual del registro o cuenta.')
bullet(doc, 'En los listados, los iconos de lápiz permiten editar; los iconos de ojo cambian el estado de un hábitat y el candado abierto rehabilita una cuenta.')
doc.add_page_break()

heading(doc, '2. Panel de cuidadores', 1)
text_para(doc, 'El panel reúne el equipo de cuidadores y sus operaciones relacionadas. En la parte superior se ven cuatro indicadores: total de cuidadores, personal en servicio, animales asignados y tareas activas.')
heading(doc, 'Consultar y localizar un cuidador', 2)
numbered(doc, 'Abra el apartado Cuidadores desde el panel de Administración.')
numbered(doc, 'Use el campo Buscar para escribir el nombre o la especialidad.')
numbered(doc, 'Seleccione Activos, Inactivos o Todos en el filtro Estado.')
numbered(doc, 'Revise la tabla, que muestra cuidador, contratación, especialidad, estado, cantidad de animales, horarios y acciones.')
heading(doc, 'Registrar un cuidador', 2)
numbered(doc, 'Seleccione + Registrar cuidador o Registrar cuidador.')
numbered(doc, 'Elija la cuenta de usuario que corresponde al cuidador.')
numbered(doc, 'Indique fecha de contratación, estado y especialidad.')
numbered(doc, 'Pulse Registrar cuidador para guardar.')
callout(doc, 'Importante.', 'El cuidador se vincula a una cuenta de usuario. Cree primero la cuenta en Usuarios si la persona aún no aparece como opción.')
heading(doc, 'Gestión individual', 2)
text_para(doc, 'Desde las acciones de cada fila se puede consultar la ficha del cuidador, editar sus datos y administrar sus relaciones operativas: animales asignados, horarios y tareas. La ficha concentra los registros asociados para facilitar su revisión.')
data_table(doc, ['Opción', 'Información que se administra'], [
    ['Animales', 'Animal asignado, fecha de asignación y responsabilidad.'],
    ['Horarios', 'Día de semana, hora de entrada y hora de salida.'],
    ['Tareas', 'Tarea del catálogo, descripción, frecuencia, hora programada y observaciones.'],
    ['Catálogo de tareas', 'Nombre, descripción y estado de las tareas disponibles para asignar.'],
], [2500,6860])
doc.add_page_break()

heading(doc, '3. Asignaciones de cuidadores', 1)
heading(doc, 'Asignar un animal', 2)
numbered(doc, 'Abra la gestión del cuidador y seleccione la opción para asignar animal.')
numbered(doc, 'Seleccione el animal en la lista disponible.')
numbered(doc, 'Indique la fecha de asignación y describa la responsabilidad del cuidador.')
numbered(doc, 'Guarde la asignación.')
text_para(doc, 'La tabla de asignaciones permite revisar el animal, la fecha, la responsabilidad, el estado y las acciones. Al finalizar o inactivar una asignación, verifique que ya no sea necesaria antes de continuar.')
heading(doc, 'Programar horario', 2)
numbered(doc, 'Abra Horario del cuidador desde su gestión individual.')
numbered(doc, 'Seleccione el día de la semana.')
numbered(doc, 'Ingrese la hora de entrada y la hora de salida.')
numbered(doc, 'Guarde y confirme que el horario aparece en el listado.')
heading(doc, 'Asignar y controlar tareas', 2)
numbered(doc, 'Abra Asignar tarea para el cuidador.')
numbered(doc, 'Elija una tarea del catálogo o registre la información solicitada: descripción, frecuencia, hora y observaciones.')
numbered(doc, 'Guarde la asignación.')
numbered(doc, 'Use las acciones de la tabla para finalizar o reactivar una tarea cuando corresponda.')
callout(doc, 'Buena práctica.', 'Mantenga activos solo los horarios, animales y tareas vigentes. Así los indicadores del panel reflejarán la operación actual.')
doc.add_page_break()

heading(doc, '4. Administración de hábitats', 1)
text_para(doc, 'Esta pantalla incluye un resumen de hábitats registrados, activos e inactivos; un formulario de alta o edición y un listado filtrable. Al editar, el mismo formulario muestra el botón Guardar cambios.')
heading(doc, 'Registrar o editar un hábitat', 2)
data_table(doc, ['Campo', 'Uso'], [
    ['Nombre *', 'Nombre identificable del espacio. Ejemplo: Bosque Tropical.'],
    ['Tipo de hábitat *', 'Clasificación ambiental. Ejemplo: Selva tropical.'],
    ['Zona *', 'Ubicación interna, como Zona Norte.'],
    ['Área (m²) *', 'Extensión del hábitat; admite decimales mayores a 0.'],
    ['Capacidad de animales *', 'Máximo de animales permitido; debe ser al menos 1.'],
    ['Icono y estado', 'Representación visual y disponibilidad: Activo o Inactivo.'],
    ['Foto', 'Imagen opcional del hábitat.'],
    ['Animales asignados', 'Seleccione los animales que pertenecerán al espacio.'],
    ['Características y descripción *', 'Condiciones, rasgos y explicación del ambiente.'],
], [2600,6760])
heading(doc, 'Pasos', 2)
numbered(doc, 'Complete los campos obligatorios y agregue foto, icono o animales cuando sean necesarios.')
numbered(doc, 'Revise la capacidad antes de guardar para asegurar que sea coherente con las asignaciones.')
numbered(doc, 'Pulse Registrar hábitat. Para modificar uno existente, pulse el icono de lápiz del listado, actualice los datos y guarde los cambios.')
callout(doc, 'Al modificar animales.', 'Si desmarca un animal que ya estaba asignado, el sistema lo deja sin hábitat. Revise esta selección cuidadosamente antes de guardar.')
doc.add_page_break()

heading(doc, '5. Listado y estado de hábitats', 1)
heading(doc, 'Buscar y filtrar', 2)
numbered(doc, 'En el listado Hábitats registrados, escriba una palabra en el buscador. Puede buscar por nombre, tipo, zona o características.')
numbered(doc, 'Elija Todos los estados, Activos o Inactivos.')
numbered(doc, 'Pulse Filtrar. Para volver a ver todo el catálogo, use Limpiar.')
heading(doc, 'Interpretar el listado', 2)
text_para(doc, 'Cada fila muestra una miniatura o icono, nombre y descripción, tipo, zona, área, capacidad y estado. Las acciones se ubican a la derecha.')
data_table(doc, ['Acción', 'Resultado'], [
    ['Lápiz', 'Abre el hábitat en modo edición dentro del formulario superior.'],
    ['Ojo tachado', 'Inactiva un hábitat activo.'],
    ['Ojo', 'Activa nuevamente un hábitat inactivo.'],
], [2200,7160])
callout(doc, 'Recomendación.', 'Inactivar es preferible a eliminar cuando se desea conservar el historial del espacio y sus datos. Antes de inactivar, compruebe el impacto sobre los animales asignados.')
doc.add_page_break()

heading(doc, '6. Usuarios y roles', 1)
text_para(doc, 'La pantalla Usuarios y roles permite administrar las cuentas que acceden al sistema. Incluye un formulario para registrar, una vista de edición y una tabla de cuentas registradas.')
heading(doc, 'Crear una cuenta', 2)
numbered(doc, 'En Nueva cuenta / Registrar usuario complete nombre completo, correo, teléfono, nacimiento y nombre de usuario.')
numbered(doc, 'Seleccione el Rol correspondiente.')
numbered(doc, 'Defina una contraseña temporal entre 8 y 72 caracteres.')
numbered(doc, 'Pulse Crear usuario.')
heading(doc, 'Editar una cuenta', 2)
numbered(doc, 'En la tabla Usuarios del sistema, pulse el icono de lápiz de la persona que desea modificar.')
numbered(doc, 'Actualice los campos requeridos. La contraseña nueva es opcional: déjela vacía si no se desea cambiar.')
numbered(doc, 'Pulse Guardar cambios.')
heading(doc, 'Estados de cuenta y rehabilitación', 2)
text_para(doc, 'La tabla muestra persona, usuario, rol, estado y acciones. Una cuenta puede figurar como Activa o Bloqueada. Si la cuenta está bloqueada o tiene intentos fallidos, aparece el botón de desbloqueo.')
numbered(doc, 'Pulse el icono de candado abierto de la cuenta afectada.')
numbered(doc, 'Confirme la rehabilitación en el mensaje de confirmación.')
numbered(doc, 'Compruebe que el estado vuelva a Activa y que los intentos hayan quedado resueltos.')
callout(doc, 'Seguridad.', 'Asigne un rol conforme a las responsabilidades reales de la persona. Evite compartir usuarios o reutilizar contraseñas temporales.')
doc.add_page_break()

heading(doc, '7. Roles y permisos', 1)
text_para(doc, 'Este apartado controla las funciones disponibles para cada rol. Las tarjetas superiores muestran los roles y cuántos permisos tienen asignados. Al elegir una tarjeta se abre su configuración.')
heading(doc, 'Configurar permisos de un rol', 2)
numbered(doc, 'Abra Roles y permisos desde el módulo administrativo.')
numbered(doc, 'Seleccione el rol que desea configurar.')
numbered(doc, 'Marque los permisos necesarios en las tarjetas de permisos. Cada tarjeta incluye el nombre y una breve descripción.')
numbered(doc, 'Pulse Guardar permisos.')
numbered(doc, 'Use Restablecer selección si desea descartar los cambios no guardados y recuperar lo configurado anteriormente.')
heading(doc, 'Reglas importantes', 2)
bullet(doc, 'El rol Administrador conserva acceso completo y no se puede restringir desde esta pantalla, para evitar que el sistema quede sin una cuenta de control.')
bullet(doc, 'Debe seleccionarse al menos un permiso para guardar la configuración de un rol no administrativo.')
bullet(doc, 'Mientras un rol no tenga permisos asignados, conserva sus accesos anteriores. Una vez guardada la selección, esta se aplica a los módulos protegidos.')
callout(doc, 'Control de cambios.', 'Revise el perfil de trabajo de cada rol antes de guardar. Otorgue únicamente los permisos que necesita para sus labores y vuelva a probar el acceso con una cuenta de ese rol.')
heading(doc, '8. Lista de verificación final', 1)
bullet(doc, '¿El cuidador está vinculado a la cuenta correcta y tiene asignaciones vigentes?')
bullet(doc, '¿El hábitat tiene capacidad, estado y animales asignados correctamente?')
bullet(doc, '¿La cuenta de usuario tiene el rol correcto y está activa?')
bullet(doc, '¿Los permisos otorgados son los mínimos necesarios para la función?')
text_para(doc, 'Fin del manual.', before=14, after=0, size=11, bold=True, color=GREEN, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.core_properties.title = 'Manual de usuario EcoFauna'
doc.core_properties.subject = 'Cuidadores, hábitats, usuarios y permisos'
doc.core_properties.author = 'EcoFauna'
doc.save(OUT)
print(OUT)
